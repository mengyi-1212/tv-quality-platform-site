# -*- coding: utf-8 -*-
from __future__ import annotations

from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


FONT = "Microsoft YaHei"
NAVY = RGBColor(24, 34, 55)
BLUE = RGBColor(31, 94, 183)
CYAN = RGBColor(23, 151, 177)
RED = RGBColor(202, 61, 67)
ORANGE = RGBColor(234, 135, 55)
GREEN = RGBColor(51, 143, 92)
LIGHT = RGBColor(246, 248, 252)
LINE = RGBColor(216, 222, 232)
TEXT = RGBColor(36, 44, 58)
MUTED = RGBColor(105, 117, 134)
WHITE = RGBColor(255, 255, 255)


def app_paths() -> tuple[Path, Path, Path]:
    app = Path(__file__).resolve().parents[1]
    workspace = app.parent
    desktop = workspace.parent
    return app, desktop / "报告", app / "data" / "output-reports"


def extract_text(slide) -> list[str]:
    values: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
            if text:
                values.append(text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        values.append(text)
    return values


def pick_slide_title(texts: list[str]) -> str:
    for text in texts:
        first = " ".join(text.split())
        if 2 <= len(first) <= 80:
            return first
    return ""


def analyze_reports(report_dir: Path) -> dict:
    files = sorted(path for path in report_dir.glob("*.pptx") if not path.name.startswith("._"))
    stats = []
    modules = Counter()
    for path in files:
        prs = Presentation(path)
        titles = []
        for slide in prs.slides:
            texts = extract_text(slide)
            all_text = " ".join(texts)
            title = pick_slide_title(texts)
            if title:
                titles.append(title)
            if any(word in all_text for word in ("客观数据", "亮度", "色域", "色准", "色偏", "峰值")):
                modules["客观数据"] += 1
            if any(word in all_text for word in ("RGB", "色域", "DCI", "BT2020", "NTSC")):
                modules["RGB色域"] += 1
            if any(word in all_text for word in ("主观", "问题点", "视效", "控光", "LD", "MEMC")):
                modules["主观问题"] += 1
            if any(word in all_text for word in ("结论", "总结", "建议")):
                modules["结论建议"] += 1
        stats.append(
            {
                "name": path.name,
                "slides": len(prs.slides),
                "size": path.stat().st_size,
                "titles": titles[:8],
                "width": prs.slide_width,
                "height": prs.slide_height,
            }
        )
    total_slides = sum(item["slides"] for item in stats)
    return {
        "files": stats,
        "file_count": len(stats),
        "total_slides": total_slides,
        "avg_slides": round(total_slides / len(stats), 1) if stats else 0,
        "sizes": sorted({(item["width"], item["height"]) for item in stats}),
        "modules": modules,
    }


def set_font(paragraph, size=16, color=TEXT, bold=False):
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=LINE, width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def add_text(slide, text, x, y, w, h, size=16, color=TEXT, bold=False, align=None, valign=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    if valign is not None:
        tf.vertical_anchor = valign
    lines = text.split("\n")
    for index, value in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = value
        if align is not None:
            paragraph.alignment = align
        set_font(paragraph, size=size, color=color, bold=bold)
    return box


def add_band(slide, title, eyebrow=None, dark=False):
    bg = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
    )
    fill(bg, NAVY if dark else LIGHT)
    bg.line.fill.background()
    if eyebrow:
        add_text(slide, eyebrow, Inches(0.55), Inches(0.38), Inches(2.4), Inches(0.28), 10, CYAN if dark else BLUE, True)
    add_text(slide, title, Inches(0.55), Inches(0.68), Inches(8.5), Inches(0.5), 20, WHITE if dark else TEXT, True)
    footer_color = RGBColor(189, 199, 216) if dark else MUTED
    add_text(slide, "TV Picture Quality Competitive Analysis Template", Inches(0.55), Inches(7.05), Inches(4.4), Inches(0.25), 8, footer_color)


def add_header(slide, title, section="竞品分析模板"):
    add_band(slide, title)
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(10.45), Inches(0.62), Inches(2.22), Inches(0.36))
    fill(tag, BLUE)
    tag.line.fill.background()
    add_text(slide, section, Inches(10.52), Inches(0.69), Inches(2.08), Inches(0.2), 8, WHITE, True, PP_ALIGN.CENTER)


def add_card(slide, x, y, w, h, title, body="", color=WHITE, accent=BLUE):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    fill(card, color)
    line(card, LINE, 0.7)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, Inches(0.08), h)
    fill(bar, accent)
    bar.line.fill.background()
    add_text(slide, title, x + Inches(0.18), y + Inches(0.14), w - Inches(0.32), Inches(0.28), 12, TEXT, True)
    if body:
        add_text(slide, body, x + Inches(0.18), y + Inches(0.48), w - Inches(0.32), h - Inches(0.62), 9, MUTED)
    return card


def add_placeholder(slide, x, y, w, h, title, hint="粘贴截图 / 数据图表"):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    fill(shape, RGBColor(235, 240, 248))
    line(shape, RGBColor(191, 202, 220), 1)
    add_text(slide, title, x + Inches(0.15), y + Inches(0.14), w - Inches(0.3), Inches(0.3), 11, TEXT, True, PP_ALIGN.CENTER)
    add_text(slide, hint, x + Inches(0.22), y + h / 2 - Inches(0.18), w - Inches(0.44), Inches(0.36), 10, MUTED, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return shape


def set_cell(cell, text, bg=WHITE, color=TEXT, bold=False, size=8.5, align=PP_ALIGN.CENTER):
    cell.text = text
    cell.fill.solid()
    cell.fill.fore_color.rgb = bg
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = align
        set_font(paragraph, size=size, color=color, bold=bold)


def add_table(slide, x, y, w, h, rows, cols, headers=None, widths=None):
    shape = slide.shapes.add_table(rows, cols, x, y, w, h)
    table = shape.table
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = width
    if headers:
        for col, header in enumerate(headers):
            set_cell(table.cell(0, col), header, bg=NAVY, color=WHITE, bold=True, size=8)
    for row in range(1 if headers else 0, rows):
        bg = WHITE if row % 2 else RGBColor(249, 251, 254)
        for col in range(cols):
            set_cell(table.cell(row, col), "", bg=bg, size=8)
    return table


def add_section_slide(prs, part, title, bullets, accent=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_band(slide, "", dark=True)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.0), Inches(0), Inches(0.18), Inches(7.5))
    fill(stripe, accent)
    stripe.line.fill.background()
    add_text(slide, part, Inches(0.72), Inches(2.35), Inches(2.2), Inches(0.36), 16, CYAN, True)
    add_text(slide, title, Inches(0.72), Inches(2.85), Inches(6.2), Inches(0.72), 28, WHITE, True)
    add_text(slide, bullets, Inches(0.78), Inches(4.0), Inches(7.0), Inches(1.2), 13, RGBColor(214, 222, 236))
    return slide


def create_template(analysis: dict, output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000

    # 1 Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_band(slide, "", dark=True)
    fill(slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.16), Inches(7.5)), RED)
    add_text(slide, "竞品分析报告", Inches(0.78), Inches(2.18), Inches(7.6), Inches(0.72), 34, WHITE, True)
    add_text(slide, "{品牌型号A}  vs  {品牌型号B}", Inches(0.82), Inches(3.02), Inches(6.4), Inches(0.38), 17, RGBColor(217, 225, 239))
    add_text(slide, "客观数据 / 主观视效 / 问题点 / 调试建议", Inches(0.82), Inches(3.58), Inches(6.8), Inches(0.32), 13, CYAN)
    add_card(slide, Inches(8.2), Inches(2.15), Inches(3.7), Inches(1.85), "报告信息", "项目：{项目名称}\n日期：{YYYY.MM.DD}\n负责人：{姓名}\n版本：V1.0", RGBColor(31, 43, 68), CYAN)
    add_text(slide, "模板来源：桌面/报告 文件夹内竞品分析 PPT 的共性结构", Inches(0.82), Inches(6.75), Inches(6.8), Inches(0.25), 8, RGBColor(178, 188, 205))

    # 2 Analysis basis
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "报告结构分析与模板说明")
    metrics = [
        ("已分析报告", f"{analysis['file_count']} 份"),
        ("总页数", f"{analysis['total_slides']} 页"),
        ("平均页数", f"{analysis['avg_slides']} 页/份"),
        ("页面比例", "16:9 宽屏"),
    ]
    for index, (title, value) in enumerate(metrics):
        add_card(slide, Inches(0.7 + index * 3.05), Inches(1.45), Inches(2.65), Inches(0.95), title, value, WHITE, [BLUE, CYAN, ORANGE, GREEN][index])
        add_text(slide, value, Inches(0.95 + index * 3.05), Inches(1.95), Inches(2.1), Inches(0.28), 17, TEXT, True, PP_ALIGN.CENTER)
    add_card(slide, Inches(0.7), Inches(2.78), Inches(5.75), Inches(2.95), "统一结构", "1. 封面与测试对象\n2. 客观数据总览\n3. 亮度/功率/峰值细节\n4. RGB色域、色准、色偏、屏对比度\n5. 主观问题点与专项问题\n6. 结论、调试建议、风险优先级", WHITE, BLUE)
    source_table = add_table(
        slide,
        Inches(6.72),
        Inches(2.78),
        Inches(5.85),
        Inches(2.95),
        rows=min(7, analysis["file_count"] + 1),
        cols=3,
        headers=["源报告", "页数", "代表模块"],
        widths=[Inches(3.25), Inches(0.75), Inches(1.85)],
    )
    for row_index, item in enumerate(analysis["files"][:6], start=1):
        title_hint = " / ".join(item["titles"][1:3]) if len(item["titles"]) > 2 else "客观/主观/结论"
        set_cell(source_table.cell(row_index, 0), item["name"], bg=WHITE if row_index % 2 else RGBColor(249, 251, 254), size=7.2, align=PP_ALIGN.LEFT)
        set_cell(source_table.cell(row_index, 1), str(item["slides"]), bg=WHITE if row_index % 2 else RGBColor(249, 251, 254), size=8.2)
        set_cell(source_table.cell(row_index, 2), title_hint[:24], bg=WHITE if row_index % 2 else RGBColor(249, 251, 254), size=7.2, align=PP_ALIGN.LEFT)

    add_section_slide(prs, "PART 01", "客观数据分析", "数据总览 -> 细节项对比 -> 关键差异解释\n建议每页只放一个判断主题，结论句放在右侧或底部。", BLUE)

    # 4 Objective overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "客观数据总览", "PART 01")
    add_text(slide, "结论摘要：{一句话说明本机相对竞品的客观优势/短板}", Inches(0.72), Inches(1.28), Inches(11.6), Inches(0.35), 13, RED, True)
    table = add_table(
        slide,
        Inches(0.72),
        Inches(1.85),
        Inches(11.9),
        Inches(4.55),
        rows=9,
        cols=7,
        headers=["维度", "测试项", "本机", "竞品A", "竞品B", "差异判断", "优先级"],
        widths=[Inches(1.08), Inches(2.08), Inches(1.28), Inches(1.28), Inches(1.28), Inches(3.05), Inches(0.85)],
    )
    rows = [
        ("亮度", "全白/峰值/Real Scene", "", "", "", "填入差异、原因和用户影响", "P0/P1"),
        ("功率", "HDR/SDR典型功耗", "", "", "", "标注亮度效率是否占优", "P1"),
        ("色域", "BT.2020/DCI-P3/NTSC", "", "", "", "说明覆盖率、四色波长策略", "P0/P1"),
        ("色准", "ΔE Avg/Max", "", "", "", "判断默认模式是否可交付", "P0"),
        ("色偏", "白平衡/RGB Balance", "", "", "", "标记偏红/偏绿/偏蓝趋势", "P1"),
        ("对比度", "ANSI/暗场/LD", "", "", "", "判断暗场黑位与泛光", "P0/P1"),
        ("均匀性", "亮度/色温均匀性", "", "", "", "说明边角、脏屏、色斑风险", "P2"),
        ("结论", "综合评级", "", "", "", "A/B/C 评级 + 调试方向", "P0"),
    ]
    for row_index, row_values in enumerate(rows, start=1):
        for col_index, value in enumerate(row_values):
            set_cell(table.cell(row_index, col_index), value, bg=WHITE if row_index % 2 else RGBColor(249, 251, 254), size=7.5, align=PP_ALIGN.LEFT if col_index in {1, 5} else PP_ALIGN.CENTER)

    # 5 Brightness
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "客观数据细节项对比 | 亮度 / 功率 / 峰值", "PART 01")
    add_placeholder(slide, Inches(0.72), Inches(1.45), Inches(7.15), Inches(4.48), "亮度与功率曲线")
    add_card(slide, Inches(8.15), Inches(1.45), Inches(4.15), Inches(1.12), "关键结论", "{本机在 HDR 峰值 / Real Scene / 功耗效率上的表现}", WHITE, RED)
    add_card(slide, Inches(8.15), Inches(2.78), Inches(4.15), Inches(1.12), "差异原因", "{背光分区、ABL、热保护、算法策略}", WHITE, ORANGE)
    add_card(slide, Inches(8.15), Inches(4.11), Inches(4.15), Inches(1.12), "调试建议", "{优先改动项、风险和回归验证指标}", WHITE, GREEN)
    add_text(slide, "推荐图表：10%窗口峰值、全白亮度、Real Scene、功耗效率，统一单位和测试模式。", Inches(0.85), Inches(6.25), Inches(11.2), Inches(0.3), 9, MUTED)

    # 6 Gamut
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "客观数据 | RGB 色域 / 四色 B-C 波长", "PART 01")
    add_placeholder(slide, Inches(0.72), Inches(1.42), Inches(5.62), Inches(4.7), "CIE 1931 / CIE 1976 色域图")
    table = add_table(
        slide,
        Inches(6.65),
        Inches(1.42),
        Inches(5.67),
        Inches(2.58),
        rows=5,
        cols=5,
        headers=["项目", "本机", "竞品A", "竞品B", "判断"],
        widths=[Inches(1.2), Inches(1.05), Inches(1.05), Inches(1.05), Inches(1.32)],
    )
    gamut_rows = [
        ("R波长", "643nm", "", "", "固定基准"),
        ("G波长", "524nm", "", "", "固定基准"),
        ("B波长", "{nm}", "", "", "蓝色覆盖"),
        ("C波长", "{nm}", "", "", "青色补偿"),
    ]
    for row_index, row_values in enumerate(gamut_rows, start=1):
        for col_index, value in enumerate(row_values):
            set_cell(table.cell(row_index, col_index), value, bg=WHITE if row_index % 2 else RGBColor(249, 251, 254), size=7.8)
    add_card(slide, Inches(6.65), Inches(4.25), Inches(2.72), Inches(1.32), "应用判定", "按目标色域、亮度效率、肤色/天空/草地场景判断 B/C 调用比例。", WHITE, BLUE)
    add_card(slide, Inches(9.62), Inches(4.25), Inches(2.7), Inches(1.32), "风险项", "关注蓝青断层、白点漂移、色准 ΔE 和低亮色偏。", WHITE, RED)

    # 7 Accuracy
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "客观数据 | 色准 / 色偏 / 屏对比度", "PART 01")
    add_placeholder(slide, Inches(0.72), Inches(1.45), Inches(3.65), Inches(3.38), "色准 ΔE 图")
    add_placeholder(slide, Inches(4.82), Inches(1.45), Inches(3.65), Inches(3.38), "RGB Balance / 色温")
    add_placeholder(slide, Inches(8.92), Inches(1.45), Inches(3.4), Inches(3.38), "对比度 / 暗场图")
    add_card(slide, Inches(0.72), Inches(5.05), Inches(3.65), Inches(0.88), "色准结论", "{平均/最大 ΔE 是否达标}", WHITE, BLUE)
    add_card(slide, Inches(4.82), Inches(5.05), Inches(3.65), Inches(0.88), "色偏结论", "{偏色方向与影响场景}", WHITE, ORANGE)
    add_card(slide, Inches(8.92), Inches(5.05), Inches(3.4), Inches(0.88), "对比度结论", "{黑位/泛光/控光策略}", WHITE, GREEN)

    add_section_slide(prs, "PART 02", "主观视效与问题点", "问题截图 -> 复现条件 -> 根因判断 -> 调试建议\n主观页需和前面的客观指标互相印证。", RED)

    # 9 Subjective overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "主观视效问题点总览", "PART 02")
    table = add_table(
        slide,
        Inches(0.72),
        Inches(1.5),
        Inches(11.85),
        Inches(4.75),
        rows=8,
        cols=6,
        headers=["问题", "场景", "本机表现", "竞品对比", "可能原因", "优先级"],
        widths=[Inches(1.75), Inches(1.75), Inches(2.15), Inches(2.15), Inches(3.0), Inches(0.9)],
    )
    issue_rows = [
        ("暗场泛光/黑位", "电影暗场", "", "", "LD/黑位补偿/伽马", "P0"),
        ("肤色偏色", "人物特写", "", "", "白平衡/色彩管理", "P0"),
        ("高光压缩", "HDR高亮", "", "", "Tone Mapping/ABL", "P1"),
        ("运动拖影", "体育/游戏", "", "", "MEMC/响应时间", "P1"),
        ("色彩断层", "天空/渐变", "", "", "B/C调用/量化/降噪", "P1"),
        ("艺术模式", "静态图片", "", "", "色温/锐化/环境光", "P2"),
        ("其他", "{补充}", "", "", "", "P2"),
    ]
    for row_index, row_values in enumerate(issue_rows, start=1):
        for col_index, value in enumerate(row_values):
            set_cell(table.cell(row_index, col_index), value, bg=WHITE if row_index % 2 else RGBColor(249, 251, 254), size=7.4, align=PP_ALIGN.LEFT if col_index < 5 else PP_ALIGN.CENTER)

    # 10 LD
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "专项问题分析 | 控光 / LD / 暗场", "PART 02")
    add_placeholder(slide, Inches(0.72), Inches(1.48), Inches(5.62), Inches(3.32), "本机截图")
    add_placeholder(slide, Inches(6.7), Inches(1.48), Inches(5.62), Inches(3.32), "竞品截图")
    add_card(slide, Inches(0.72), Inches(5.05), Inches(3.65), Inches(0.9), "复现条件", "{片源 / 模式 / 亮度 / LD档位}", WHITE, BLUE)
    add_card(slide, Inches(4.82), Inches(5.05), Inches(3.65), Inches(0.9), "根因判断", "{分区响应 / 黑位策略 / 伽马}", WHITE, ORANGE)
    add_card(slide, Inches(8.92), Inches(5.05), Inches(3.4), Inches(0.9), "建议", "{参数方向与验证点}", WHITE, GREEN)

    # 11 MEMC/art
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "专项问题分析 | MEMC / 运动 / 艺术模式", "PART 02")
    add_placeholder(slide, Inches(0.72), Inches(1.48), Inches(3.65), Inches(2.75), "运动场景截图")
    add_placeholder(slide, Inches(4.82), Inches(1.48), Inches(3.65), Inches(2.75), "艺术模式截图")
    add_placeholder(slide, Inches(8.92), Inches(1.48), Inches(3.4), Inches(2.75), "竞品对照")
    add_card(slide, Inches(0.72), Inches(4.55), Inches(3.65), Inches(1.28), "MEMC判断", "拖影、抖动、肥皂感、插帧错误。", WHITE, BLUE)
    add_card(slide, Inches(4.82), Inches(4.55), Inches(3.65), Inches(1.28), "艺术模式判断", "色温、亮度、锐化、环境光适配。", WHITE, ORANGE)
    add_card(slide, Inches(8.92), Inches(4.55), Inches(3.4), Inches(1.28), "回归验证", "同片源、同模式、同环境复测。", WHITE, GREEN)

    # 12 Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "结论与调试建议", "结论")
    add_card(slide, Inches(0.72), Inches(1.45), Inches(3.55), Inches(1.42), "综合结论", "{一句话说明本机竞争力、主要优势和必须解决的短板}", WHITE, RED)
    add_card(slide, Inches(4.72), Inches(1.45), Inches(3.55), Inches(1.42), "优先改进", "P0：{必须处理}\nP1：{建议优化}\nP2：{观察项}", WHITE, ORANGE)
    add_card(slide, Inches(8.72), Inches(1.45), Inches(3.55), Inches(1.42), "交付风险", "{量产/调试/体验风险与验收口径}", WHITE, BLUE)
    table = add_table(
        slide,
        Inches(0.72),
        Inches(3.25),
        Inches(11.85),
        Inches(2.8),
        rows=5,
        cols=5,
        headers=["优先级", "问题", "调试方向", "验证指标", "负责人/日期"],
        widths=[Inches(1.0), Inches(2.4), Inches(3.1), Inches(3.0), Inches(2.35)],
    )
    action_rows = [
        ("P0", "{核心问题}", "{参数/算法/模式}", "{客观指标+主观场景}", "{姓名/日期}"),
        ("P1", "{优化项}", "{参数/算法/模式}", "{复测标准}", "{姓名/日期}"),
        ("P2", "{观察项}", "{后续跟踪}", "{量产抽检}", "{姓名/日期}"),
        ("关闭", "{已解决项}", "{固化版本}", "{回归通过}", "{版本号}"),
    ]
    for row_index, row_values in enumerate(action_rows, start=1):
        for col_index, value in enumerate(row_values):
            set_cell(table.cell(row_index, col_index), value, bg=WHITE if row_index % 2 else RGBColor(249, 251, 254), size=7.8, align=PP_ALIGN.LEFT if col_index > 0 else PP_ALIGN.CENTER)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def main() -> None:
    _, report_dir, output_dir = app_paths()
    if not report_dir.exists():
        raise SystemExit(f"未找到报告文件夹：{report_dir}")
    analysis = analyze_reports(report_dir)
    output_path = output_dir / "输出报告-竞品分析PPT模板.pptx"
    create_template(analysis, output_path)
    print(f"已生成模板：{output_path}")
    print(f"已分析 {analysis['file_count']} 份 PPT，合计 {analysis['total_slides']} 页。")


if __name__ == "__main__":
    main()
